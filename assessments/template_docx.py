"""
assessments.template_docx
=========================

Builds the downloadable Word templates trainers fill in and upload. They are
written in exactly the layout importer.py reads, and a test round-trips both
through the importer so the two can never drift apart.

Lines beginning with "//" are guidance notes: the importer ignores them, so
they can stay in the file or be deleted.
"""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

GREY = RGBColor(0x6B, 0x72, 0x80)

NOTES = [
    "// How to use this template: replace the text, keep the layout. Lines starting with // are notes and are ignored when you upload.",
    "// Number each question (1. 2. 3.). Put options on lines starting A. B. C. D. and the correct one on an Answer: line.",
    "// Write the marks at the end of a question, for example (2 marks). Delete any section you don't need.",
    "// More than one correct option? Write Answer: A, C . Candidates then tick all that apply (you can change that on the review screen).",
]


def _note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    return p


def _p(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def build(kind):
    """kind: 'blank' (placeholders to overwrite) or 'example' (a finished sample)."""
    ex = kind == "example"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = "Assessment import template. Keep the layout; the system reads it when you upload."
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = GREY

    title = doc.add_paragraph(style="Title")
    title.add_run("Networking Fundamentals Test" if ex else "Assessment title")
    _p(doc, "Answer all questions. Time allowed: 1 hour." if ex else "Instructions for candidates (optional): for example, time allowed and rules.")
    for n in NOTES:
        _note(doc, n)

    # ---- A: multiple choice --------------------------------------------------
    doc.add_heading("Section A: Multiple Choice (4 marks)", level=1)
    _note(doc, "// One correct option per question. Instead of an Answer: line you may make the correct option bold, or list answers in an Answer Key at the end.")
    mcq = [
        ("Which OSI layer is responsible for routing packets?", ["Physical", "Network", "Session", "Application"], "B")
        if ex else ("Type the first question here.", ["First option", "Second option", "Third option", "Fourth option"], "A"),
        ("What is the default port for HTTP?", ["21", "80", "443", "3306"], "B")
        if ex else ("Type the second question here.", ["First option", "Second option", "Third option", "Fourth option"], "A"),
    ]
    for i, (q, opts, ans) in enumerate(mcq, 1):
        _p(doc, f"{i}. {q} (2 marks)")
        for letter, o in zip("ABCD", opts):
            _p(doc, f"{letter}. {o}")
        _p(doc, f"Answer: {ans}")

    # ---- B: true / false ----------------------------------------------------------
    doc.add_heading("Section B: True or False (2 marks)", level=1)
    _note(doc, "// Write the statement, then Answer: True or Answer: False.")
    tf = [("A switch connects devices within the same network.", "True"),
          ("An IP address identifies a device on a network.", "True")] if ex else \
         [("Type the first statement here.", "True"), ("Type the second statement here.", "False")]
    for i, (q, a) in enumerate(tf, 3):
        _p(doc, f"{i}. {q} (1 mark)")
        _p(doc, f"Answer: {a}")

    # ---- C: fill in the blank -----------------------------------------------------------
    doc.add_heading("Section C: Fill in the Blanks (4 marks)", level=1)
    _note(doc, "// Mark the gap with ____ . Give every accepted answer on the Answer: line, separated by / . Capital letters and extra spaces are ignored.")
    fills = [("The ____ protocol translates domain names into IP addresses.", "DNS / domain name system"),
             ("A ____ connects different networks together.", "router")] if ex else \
            [("Type a sentence with ____ where the answer goes.", "accepted answer / another accepted answer"),
             ("Type another sentence with ____ here.", "accepted answer")]
    for i, (q, a) in enumerate(fills, 5):
        _p(doc, f"{i}. {q} (2 marks)")
        _p(doc, f"Answer: {a}")

    # ---- D: matching ---------------------------------------------------------------------
    doc.add_heading("Section D: Matching (3 marks)", level=1)
    _note(doc, "// Number the items to match and letter the choices. Give the key as Answer: 1-B, 2-A, 3-C . You may add extra wrong choices (here D).")
    _p(doc, "7. Match each device in Column A with its function in Column B. (3 marks)")
    left = ["Router", "Switch", "Access point"] if ex else ["First item", "Second item", "Third item"]
    right = ["Connects separate networks", "Connects devices in one network", "Provides wireless access", "Stores web pages"] if ex else \
            ["Match for the first item", "Match for the second item", "Match for the third item", "Extra wrong choice (optional)"]
    for n, t in enumerate(left, 1):
        _p(doc, f"{n}. {t}")
    for letter, t in zip("ABCD", right):
        _p(doc, f"{letter}. {t}")
    _p(doc, "Answer: 1-A, 2-B, 3-C")

    # ---- E: open questions -------------------------------------------------------------------
    doc.add_heading("Section E: Open Questions (10 marks)", level=1)
    _note(doc, "// Marked by the trainer. The Marking guide: line is for trainers only and is never shown to candidates.")
    opens = [("Explain how DHCP assigns IP addresses to devices. (5 marks)",
              "Marking guide: Server leases an address from a pool; DORA exchange; lease time and renewal.")
             if ex else ("Type the first open question here. (5 marks)", "Marking guide: What a full-mark answer should contain.")]
    opens.append(("Describe two differences between a hub and a switch. (5 marks)",
                  "Marking guide: Switch forwards by MAC address; hub broadcasts to all ports; collision domains.")
                 if ex else ("Type the second open question here. (5 marks)", "Marking guide: What a full-mark answer should contain."))
    for i, (q, g) in enumerate(opens, 8):
        _p(doc, f"{i}. {q}")
        _p(doc, g)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
