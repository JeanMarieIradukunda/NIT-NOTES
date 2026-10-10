"""
assessments.trainer_tools
=========================

Four trainer conveniences built on top of the staff views:

* ``marking_queue``    every submitted attempt that still has unmarked written answers
* ``guides_edit``      write or fix every written-answer marking guide on one page
* ``guide_download``   the marking guide as a Word file or a PDF
* ``exam_duplicate``   copy an assessment (settings, questions, answer keys) as a closed draft
"""

import io
import re
import secrets

from django.contrib import messages
from django.db import transaction
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from accounts.roles import is_admin

from .marking import question_max_marks
from .models import AnswerKey, Attempt, Exam, Question, SECTION_ORDER
from .permissions import trainer_required
from .staff_views import GUIDE_LABELS, _exam_or_deny, _guide_row

GUIDE_MAX_CHARS = 5000


# --------------------------------------------------------------------------- #
# Marking queue
# --------------------------------------------------------------------------- #

@trainer_required
def marking_queue(request):
    """Submitted attempts whose written answers are not marked yet, oldest first."""
    attempts = (Attempt.objects.filter(status=Attempt.SUBMITTED, marking_complete=False)
                .select_related("exam", "exam__module").order_by("submitted_at", "id"))
    if not is_admin(request.user):
        attempts = attempts.filter(exam__created_by=request.user)
    rows = list(attempts)
    now = timezone.now()
    for a in rows:
        a.waiting_days = (now - a.submitted_at).days if a.submitted_at else 0
    return render(request, "assessments/marking_queue.html", {
        "attempts": rows, "show_owner": is_admin(request.user)})


# --------------------------------------------------------------------------- #
# Edit every guide on one page
# --------------------------------------------------------------------------- #

@trainer_required
def guides_edit(request, pk):
    exam, denied = _exam_or_deny(request, pk)
    if denied:
        return denied
    questions = list(exam.questions.filter(section="open").select_related("key"))
    maxima = question_max_marks(exam, list(exam.questions.all()))

    def current(q):
        try:
            return (q.key.data or {}).get("guide", "")
        except AnswerKey.DoesNotExist:
            return ""

    if request.method == "POST":
        new = {q.pk: (request.POST.get(f"guide_{q.pk}") or "").strip() for q in questions}
        too_long = [q.order for q in questions if len(new[q.pk]) > GUIDE_MAX_CHARS]
        if too_long:
            messages.error(request, f"A guide can be at most {GUIDE_MAX_CHARS} characters. "
                                    f"Shorten question {', '.join(map(str, too_long))} and save again. "
                                    "Nothing was saved.")
        else:
            changed = 0
            with transaction.atomic():
                for q in questions:
                    if new[q.pk] == (current(q) or "").strip():
                        continue
                    key, _ = AnswerKey.objects.get_or_create(question=q, defaults={"data": {}})
                    data = dict(key.data or {})
                    data["guide"] = new[q.pk]
                    key.data = data
                    key.save(update_fields=["data"])
                    changed += 1
            messages.success(request, f"Saved {changed} guide{'s' if changed != 1 else ''}."
                             if changed else "No changes to save.")
            return redirect("assessments:exam_guide", pk=exam.pk)
        rows = [{"q": q, "marks": maxima[q.id], "guide": new[q.pk]} for q in questions]
    else:
        rows = [{"q": q, "marks": maxima[q.id], "guide": current(q)} for q in questions]
    return render(request, "assessments/guides_edit.html", {
        "exam": exam, "rows": rows, "max_chars": GUIDE_MAX_CHARS})


# --------------------------------------------------------------------------- #
# Download (Word / PDF)
# --------------------------------------------------------------------------- #

def _answer_lines(key, r):
    """The answer for one question as plain lines (shared by the Word and PDF versions)."""
    if key == "mcq":
        if r["tf"]:
            right = [o["text"] for o in r["options"] if o["correct"]]
            return [f"Answer: {', '.join(right)}" if right else "Answer: not set"]
        lines = [f"{o['letter']}. {o['text']}" + ("   ✓ correct" if o["correct"] else "") for o in r["options"]]
        if r["multi"]:
            lines.append("(select all that apply)")
        elif r["accept_any"]:
            lines.append("(any one of the ticked answers is accepted)")
        return lines
    if key == "fill":
        tail = " (case sensitive)" if r["case_sensitive"] else ""
        return ["Accepted: " + " · ".join(r["accepted"]) + tail if r["accepted"] else "Accepted: not set"]
    if key == "match":
        return [f"{left}  →  {right}" for left, right in r["pairs"]]
    return (r["guide"].splitlines() if r["guide"] else ["(no marking guide yet)"])


def _plain_guide(exam):
    questions = list(exam.questions.select_related("key"))
    maxima = question_max_marks(exam, questions)
    sections = []
    for s in SECTION_ORDER:
        rows = []
        for q in (x for x in questions if x.section == s):
            r = _guide_row(q, maxima[q.id])
            rows.append({"number": q.order, "text": q.text, "marks": maxima[q.id],
                         "lines": _answer_lines(s, r)})
        if rows:
            sections.append({"label": GUIDE_LABELS[s], "rows": rows})
    return sections


def _filename(exam, ext):
    slug = re.sub(r"[^A-Za-z0-9]+", "-", exam.title).strip("-").lower()[:50] or "assessment"
    return f"marking-guide-{slug}.{ext}"


def _word(exam, sections):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(f"Marking guide: {exam.title}", level=1)
    doc.add_paragraph(f"{exam.total_marks:g} marks · {exam.duration_minutes} minutes · for trainers only")
    for sec in sections:
        doc.add_heading(sec["label"], level=2)
        for r in sec["rows"]:
            p = doc.add_paragraph()
            p.add_run(f"{r['number']}. {r['text']}").bold = True
            p.add_run(f"   [{r['marks']:g} mark{'s' if r['marks'] != 1 else ''}]")
            for line in r["lines"]:
                doc.add_paragraph(line, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@trainer_required
def guide_download(request, pk, fmt):
    exam, denied = _exam_or_deny(request, pk)
    if denied:
        return denied
    if fmt not in ("docx", "pdf"):
        raise Http404("Unknown format")
    sections = _plain_guide(exam)
    if fmt == "docx":
        resp = HttpResponse(_word(exam, sections), content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        resp["Content-Disposition"] = f'attachment; filename="{_filename(exam, "docx")}"'
        return resp
    try:
        from weasyprint import HTML
        html = render_to_string("assessments/exam_guide_pdf.html", {"exam": exam, "sections": sections})
        pdf = HTML(string=html).write_pdf()
    except (ImportError, OSError):
        messages.error(request, "PDF export is not available on this server. Use the Word download "
                                "or the Print button instead.")
        return redirect("assessments:exam_guide", pk=exam.pk)
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{_filename(exam, "pdf")}"'
    return resp


# --------------------------------------------------------------------------- #
# Copy an assessment
# --------------------------------------------------------------------------- #

COPIED_FIELDS = (
    "module", "instructions", "max_opens", "max_devices", "duration_minutes", "max_violations",
    "penalties", "marks_mcq", "marks_fill", "marks_open", "marks_match", "show_results",
    "show_on_dashboard", "shuffle_questions", "shuffle_options", "multi_scoring",
)


@require_POST
@trainer_required
def exam_duplicate(request, pk):
    """A closed copy with a new password: settings, questions and answer keys come across.
    Dates, the class list, candidates and results do not."""
    source, denied = _exam_or_deny(request, pk)
    if denied:
        return denied
    password = secrets.token_urlsafe(6).replace("-", "x").replace("_", "y")
    with transaction.atomic():
        copy = Exam(title=f"Copy of {source.title}"[:200], created_by=request.user,
                    is_open=False, show_answers=False,
                    **{f: getattr(source, f) for f in COPIED_FIELDS})
        copy.set_password(password)
        copy.save()
        for q in source.questions.select_related("key"):
            new_q = Question.objects.create(exam=copy, section=q.section, order=q.order, text=q.text,
                                            weight=q.weight, payload=q.payload)
            try:
                AnswerKey.objects.create(question=new_q, data=q.key.data)
            except AnswerKey.DoesNotExist:
                pass
    messages.success(request, f"Copy created, closed. New exam password (shown once, write it down): {password}")
    messages.info(request, "Set the date, class list and open it when you are ready.")
    return redirect("assessments:exam_detail", pk=copy.pk)
