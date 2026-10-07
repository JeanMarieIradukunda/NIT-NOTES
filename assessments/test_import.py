import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from docx import Document

from . import importer
from .models import Attempt, Exam, ImportDraft, Question
from .tests import PLAIN_STATIC

User = get_user_model()


def docx_bytes(build):
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def para(doc, text, bold=False, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold = bold
    return p


def full_doc(doc):
    doc.add_paragraph("Networking Test", style="Title")
    para(doc, "Answer all questions. Time allowed: one hour.")
    doc.add_heading("Section A: Multiple Choice (4 marks)", level=1)
    para(doc, "1. Which layer routes packets? (2 marks)")
    for t, b in (("A. Physical", False), ("B. Network", True), ("C. Session", False)):
        para(doc, t, bold=b)
    para(doc, "2. Default HTTP port? (2 marks)")
    para(doc, "A. 21\nB. 80\nC. 443")
    para(doc, "Answer: B")
    doc.add_heading("Section B: Fill in the blanks (2 marks)", level=1)
    para(doc, "3. The ........ protocol resolves names. (2 marks)")
    para(doc, "Answer: DNS / domain name system.")
    doc.add_heading("Section C: Matching (2 marks)", level=1)
    para(doc, "4. Match Column A with Column B.")
    para(doc, "1. Router")
    para(doc, "2. Switch")
    para(doc, "A. Connects hosts")
    para(doc, "B. Connects networks")
    para(doc, "C. Extra choice")
    para(doc, "Answer: 1-B, 2-A")
    doc.add_heading("Section D: Open Questions (5 marks)", level=1)
    para(doc, "5. Explain DHCP. (5 marks)")
    para(doc, "(a) Define it.")
    para(doc, "Marking guide: Mentions leases and automatic addressing.")


class ParserTests(TestCase):
    def parse(self, data, name="t.docx"):
        return importer.parse_upload(name, data)

    def test_full_docx_all_four_types(self):
        d = self.parse(docx_bytes(full_doc))
        self.assertEqual(d["title"], "Networking Test")
        self.assertIn("Time allowed", d["intro"])
        qs = d["questions"]
        self.assertEqual([q["section"] for q in qs], ["mcq", "mcq", "fill", "match", "open"])
        q1, q2, q3, q4, q5 = qs
        self.assertEqual(q1["options"], ["Physical", "Network", "Session"])
        self.assertEqual((q1["correct"], q1["key_source"]), ([1], "bold / underlined option"))
        self.assertEqual(q1["text"], "Which layer routes packets?")
        self.assertEqual(q1["marks"], 2.0)
        self.assertEqual((q2["options"], q2["correct"]), (["21", "80", "443"], [1]))   # options in one paragraph
        self.assertEqual(q3["text"], "The ____ protocol resolves names.")
        self.assertEqual(q3["accepted"], ["DNS", "domain name system"])
        self.assertEqual(q4["left"], ["Router", "Switch"])
        self.assertEqual(q4["right"], ["Connects hosts", "Connects networks", "Extra choice"])
        self.assertEqual(q4["pairs"], {"0": 1, "1": 0})
        self.assertIn("(a) Define it.", q5["text"])
        self.assertIn("leases", q5["guide"])
        self.assertEqual(d["header_marks"]["mcq"], 4.0)

    def test_answer_key_block_with_compact_and_table_forms(self):
        def build(doc):
            doc.add_heading("Section A: Multiple choice", level=1)
            for n, (q, opts) in enumerate([("First?", "A. a1\nB. b1"), ("Second?", "A. a2\nB. b2"), ("Third?", "A. a3\nB. b3")], 1):
                para(doc, f"{n}. {q}")
                para(doc, opts)
            doc.add_heading("Section B: Fill in the blank", level=1)
            para(doc, "4. Water boils at ____ degrees.")
            doc.add_heading("Answer Key", level=1)
            para(doc, "1-B 2-A, 3. B")
            t = doc.add_table(rows=2, cols=2)
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = "No", "Answer"
            t.rows[1].cells[0].text, t.rows[1].cells[1].text = "4", "100"
        qs = self.parse(docx_bytes(build))["questions"]
        self.assertEqual([q.get("correct") for q in qs[:3]], [[1], [0], [1]])
        self.assertEqual(qs[3]["accepted"], ["100"])
        self.assertEqual(qs[0]["key_source"], "answer key")

    def test_key_numbers_that_restart_per_section(self):
        def build(doc):
            doc.add_heading("Section A: Multiple choice", level=1)
            para(doc, "1. Q one?"); para(doc, "A. x\nB. y")
            doc.add_heading("Section B: Fill in the blank", level=1)
            para(doc, "1. Capital of France is ____.")
            doc.add_heading("Answer Key", level=1)
            doc.add_heading("Section A: Multiple choice", level=2)
            para(doc, "1. B")
            doc.add_heading("Section B: Fill in the blank", level=2)
            para(doc, "1. Paris")
        qs = self.parse(docx_bytes(build))["questions"]
        self.assertEqual(qs[0]["correct"], [1])
        self.assertEqual(qs[1]["accepted"], ["Paris"])

    def test_word_automatic_numbering(self):
        def build(doc):
            doc.add_heading("Multiple choice questions", level=1)
            for q in ("Alpha?", "Beta?"):
                doc.add_paragraph(q, style="List Number")
                para(doc, "A. yes\nB. no")
                para(doc, "Answer: A")
        qs = self.parse(docx_bytes(build))["questions"]
        self.assertEqual([q["text"] for q in qs], ["Alpha?", "Beta?"])
        self.assertEqual([q["correct"] for q in qs], [[0], [0]])

    def test_matching_table_with_aligned_rows_warns(self):
        def build(doc):
            doc.add_heading("Section C: Matching", level=1)
            para(doc, "1. Match each term to its meaning.")
            t = doc.add_table(rows=3, cols=2)
            for r, (a, b) in zip(t.rows, [("Column A", "Column B"), ("CPU", "Processes instructions"), ("RAM", "Temporary memory")]):
                r.cells[0].text, r.cells[1].text = a, b
        q = self.parse(docx_bytes(build))["questions"][0]
        self.assertEqual((q["left"], q["right"]), (["CPU", "RAM"], ["Processes instructions", "Temporary memory"]))
        self.assertEqual(q["pairs"], {"0": 0, "1": 1})
        self.assertTrue(any("Check the pairs" in w for w in q["warnings"]))

    def test_true_false_section(self):
        def build(doc):
            doc.add_heading("Section A: True or False", level=1)
            para(doc, "1. The sun is a star.")
            para(doc, "Answer: True")
            para(doc, "2. Fish are mammals.")
            para(doc, "Answer: F")
        qs = self.parse(docx_bytes(build))["questions"]
        self.assertEqual([(q["options"], q["correct"]) for q in qs], [(["True", "False"], [0]), (["True", "False"], [1])])

    def test_no_section_headings_infers_types(self):
        def build(doc):
            para(doc, "1. Pick one?"); para(doc, "A. x"); para(doc, "B. y"); para(doc, "Answer: A")
            para(doc, "2. HTML stands for ____ Markup Language."); para(doc, "Answer: Hyper Text")
            para(doc, "3. Discuss the OSI model.")
        qs = self.parse(docx_bytes(build))["questions"]
        self.assertEqual([q["section"] for q in qs], ["mcq", "fill", "open"])

    def test_missing_answers_are_flagged_not_guessed(self):
        def build(doc):
            para(doc, "1. Pick one?"); para(doc, "A. x"); para(doc, "B. y")
        q = self.parse(docx_bytes(build))["questions"][0]
        self.assertEqual(q["correct"], [])
        self.assertTrue(any("No correct answer" in w for w in q["warnings"]))

    def test_html_with_nested_lists_inputs_and_bold_answer(self):
        html = b"""<html><head><title>x</title><style>.a{}</style><script>var secret='LEAK-ME';</script></head><body>
        <h1>IT Quiz</h1><h2>Section A: Multiple Choice</h2>
        <ol><li>Which is a browser?
              <ol type="A"><li>Excel</li><li><strong>Firefox</strong></li><li>Word</li></ol></li>
            <li>Largest planet?<ol type="A"><li>Mars</li><li>Jupiter</li></ol><p>Answer: B</p></li></ol>
        <h2>Section B: Fill in the blank</h2>
        <p>3. The capital of Rwanda is <input type="text"> .</p><p>Answer: Kigali</p>
        <h2>Section C: Matching</h2>
        <table><tr><th>Column A</th><th>Column B</th></tr><tr><td>HTML</td><td>Structure</td></tr><tr><td>CSS</td><td>Style</td></tr></table>
        </body></html>"""
        d = self.parse(html, "q.html")
        qs = d["questions"]
        self.assertEqual([q["section"] for q in qs], ["mcq", "mcq", "fill", "match"])
        self.assertEqual((qs[0]["options"], qs[0]["correct"]), (["Excel", "Firefox", "Word"], [1]))
        self.assertEqual(qs[1]["correct"], [1])
        self.assertEqual((qs[2]["text"], qs[2]["accepted"]), ("The capital of Rwanda is ____ .", ["Kigali"]))
        self.assertEqual(qs[3]["left"], ["HTML", "CSS"])
        self.assertNotIn("LEAK-ME", str(d))

    def test_bad_files_are_refused_clearly(self):
        for name, data, text in [
            ("a.exe", b"MZ", "Word (.docx) or HTML"),
            ("a.doc", b"x", "Save As"),
            ("a.docx", b"not a zip", "not a valid Word"),
            ("a.docx", bytes.fromhex("D0CF11E0A1B11AE1") + b"0" * 20, "old Word"),
            ("a.html", b"<p>just a sentence</p>", "No questions"),
        ]:
            with self.assertRaises(importer.ImportProblem) as cm:
                importer.parse_upload(name, data)
            self.assertIn(text, str(cm.exception), name)
        with self.assertRaises(importer.ImportProblem):
            importer.parse_upload("big.docx", b"0" * (importer.MAX_UPLOAD_BYTES + 1))


@override_settings(STORAGES=PLAIN_STATIC)
class ImportFlowTests(TestCase):
    def setUp(self):
        self.trainer = User.objects.create_user("t1", password="pw12345!")
        self.trainer.groups.add(Group.objects.get_or_create(name="Trainer")[0])
        self.exam = Exam.objects.create(title="Imported", created_by=self.trainer, is_open=True)
        self.exam.set_password("pw"); self.exam.save()
        self.c = Client(); self.c.force_login(self.trainer)

    def upload(self, data=None, name="test.docx"):
        data = data or docx_bytes(full_doc)
        return self.c.post(reverse("assessments:import_upload", args=[self.exam.pk]),
                           {"file": SimpleUploadedFile(name, data)})

    def test_upload_review_import_creates_standard_questions(self):
        r = self.upload()
        self.assertEqual(r.status_code, 302)
        draft = ImportDraft.objects.get()
        page = self.c.get(r["Location"])
        self.assertContains(page, "Review 5 questions")
        post = {"mode": "append", "apply_marks": "1", "use_intro": "1"}
        for i in range(5):
            post[f"inc_{i}"] = "1"
        done = self.c.post(r["Location"], post)
        self.assertEqual(done.status_code, 302)
        self.exam.refresh_from_db()
        self.assertEqual(self.exam.questions.count(), 5)
        self.assertEqual((self.exam.marks_mcq, self.exam.marks_fill, self.exam.marks_open, self.exam.marks_match),
                         (4, 2, 5, 2))
        self.assertIn("Time allowed", self.exam.instructions)
        m = self.exam.questions.get(section="match")
        self.assertEqual(m.payload["left"], ["Router", "Switch"])
        self.assertEqual(m.payload["right"], ["Connects networks", "Connects hosts", "Extra choice"])
        self.assertEqual(m.key.data["pairs"], {"0": 0, "1": 1})
        self.assertEqual(self.exam.questions.get(section="mcq", order=1).key.data["correct"], [1])
        self.assertFalse(ImportDraft.objects.filter(pk=draft.pk).exists())

    def test_missing_answers_can_be_filled_in_on_the_review_screen(self):
        def build(doc):
            para(doc, "1. Pick one?"); para(doc, "A. x"); para(doc, "B. y")
            para(doc, "2. Colour of grass is ____.")
        r = self.upload(docx_bytes(build))
        blocked = self.c.post(r["Location"], {"inc_0": "1", "inc_1": "1"})          # nothing importable yet
        self.assertEqual(Question.objects.count(), 0)
        self.c.post(r["Location"], {"inc_0": "1", "correct_0": "1", "inc_1": "1", "accepted_1": "green / Green"})
        self.assertEqual(Question.objects.count(), 2)
        self.assertEqual(Question.objects.get(section="mcq").key.data["correct"], [1])
        self.assertEqual(Question.objects.get(section="fill").key.data["accepted"], ["green", "Green"])

    def test_replace_mode_and_its_block_once_candidates_have_attempted(self):
        self.upload(); r = self.c.post(reverse("assessments:import_upload", args=[self.exam.pk]),
                                       {"file": SimpleUploadedFile("a.docx", docx_bytes(full_doc))})
        post = {"mode": "replace", **{f"inc_{i}": "1" for i in range(5)}}
        self.c.post(r["Location"], post)
        self.assertEqual(self.exam.questions.count(), 5)
        r2 = self.upload()
        self.c.post(r2["Location"], post)
        self.assertEqual(self.exam.questions.count(), 5)                              # replaced, not doubled
        Attempt.objects.create(exam=self.exam, candidate_name="A", reg_no="1")
        r3 = self.upload()
        self.c.post(r3["Location"], post)                                             # replace is ignored: appends
        self.assertEqual(self.exam.questions.count(), 10)

    def test_imported_exam_never_leaks_the_key_to_candidates(self):
        r = self.upload()
        self.c.post(r["Location"], {f"inc_{i}": "1" for i in range(5)})
        cand = Client()
        resp = cand.post(reverse("assessments:entry", args=[self.exam.public_id]),
                         {"name": "A B", "reg_no": "X1", "password": "pw"})
        page = cand.get(resp["Location"])
        html = page.content.decode()
        self.assertNotIn("leases and automatic", html)
        self.assertNotIn("domain name system", html)
        self.assertNotIn('"correct"', html)

    def test_upload_errors_and_permissions(self):
        r = self.upload(b"junk", "x.docx")
        self.assertEqual(r.status_code, 400)
        self.assertContains(r, "not a valid Word", status_code=400)
        other = User.objects.create_user("t2", password="pw12345!")
        other.groups.add(Group.objects.get(name="Trainer"))
        oc = Client(); oc.force_login(other)
        self.assertEqual(oc.get(reverse("assessments:import_upload", args=[self.exam.pk])).status_code, 403)
        self.assertEqual(Client().get(reverse("assessments:import_upload", args=[self.exam.pk])).status_code, 302)


@override_settings(STORAGES=PLAIN_STATIC)
class TemplateTests(TestCase):
    def setUp(self):
        self.trainer = User.objects.create_user("tt", password="pw12345!")
        self.trainer.groups.add(Group.objects.get_or_create(name="Trainer")[0])
        self.c = Client(); self.c.force_login(self.trainer)

    def test_guidance_notes_are_ignored(self):
        def build(doc):
            para(doc, "// a note"); doc.add_heading("Section A: Multiple choice", level=1)
            para(doc, "// another note"); para(doc, "1. Real question?"); para(doc, "A. x\nB. y"); para(doc, "Answer: A")
        d = importer.parse_upload("a.docx", docx_bytes(build))
        self.assertEqual(len(d["questions"]), 1)
        self.assertNotIn("note", d["intro"])

    def test_example_template_imports_fully_and_correctly(self):
        from . import template_docx
        d = importer.parse_upload("t.docx", template_docx.build("example"))
        qs = d["questions"]
        self.assertEqual([q["section"] for q in qs], ["mcq", "mcq", "mcq", "mcq", "fill", "fill", "match", "open", "open"])
        self.assertTrue(all(q["warnings"] == [] for q in qs), [q["warnings"] for q in qs])
        self.assertEqual([q["correct"] for q in qs[:4]], [[1], [1], [0], [0]])                 # B, B, True, True
        self.assertEqual(qs[4]["accepted"], ["DNS", "domain name system"])
        self.assertEqual(qs[6]["pairs"], {"0": 0, "1": 1, "2": 2})
        self.assertEqual(len(qs[6]["right"]), 4)
        self.assertIn("DORA", qs[7]["guide"])
        self.assertEqual(d["title"], "Networking Fundamentals Test")
        self.assertEqual(d["intro"], "Answer all questions. Time allowed: 1 hour.")
        self.assertEqual([d["header_marks"][k] for k in ("mcq", "fill", "match", "open")], [6.0, 4.0, 3.0, 10.0])

    def test_blank_template_is_readable_and_every_question_is_ready(self):
        from . import template_docx
        d = importer.parse_upload("t.docx", template_docx.build("blank"))
        self.assertEqual(len(d["questions"]), 9)
        self.assertTrue(all(q["warnings"] == [] for q in d["questions"]))
        self.assertEqual(d["title"], "Assessment title")

    def test_download_views(self):
        for kind, name in (("blank", "assessment-template-blank.docx"), ("example", "assessment-template-example.docx")):
            r = self.c.get(reverse("assessments:import_template", args=[kind]))
            self.assertEqual(r.status_code, 200)
            self.assertIn(name, r["Content-Disposition"])
            self.assertTrue(r["Content-Type"].endswith("wordprocessingml.document"))
            self.assertEqual(r.content[:2], b"PK")
        self.assertEqual(self.c.get(reverse("assessments:import_template", args=["html"])).status_code, 404)
        self.assertEqual(Client().get(reverse("assessments:import_template", args=["blank"])).status_code, 302)
        exam = Exam.objects.create(title="E", created_by=self.trainer)
        page = self.c.get(reverse("assessments:import_upload", args=[exam.pk]))
        self.assertContains(page, "Blank template (.docx)")
        self.assertNotContains(page, ".html</a>")
